"""Agrega al Excel las reseñas de Google de Angie's Secret.

Uso: python3 actualizar_resenas.py "Reseñas Google - Angie's Secret.xlsx" resenas.json AAAA-MM-DDTHH:MM [completa]

resenas.json es la lista leída de Google Maps (pestaña Reseñas, ordenadas por "Más recientes"):
[{"id": "...", "nombre": "...", "estrellas": 5, "fecha": "hace un día", "texto": "...",
  "respuesta": "...", "perfil": "2 reseñas"}]
AAAA-MM-DDTHH:MM es el día y la hora de Chile en que se leyeron (para calcular la fecha aproximada desde "hace 3 días").
"completa" indica que se leyeron TODAS las reseñas: así se puede marcar las que ya no aparecen en Google.

- Si el Excel no existe, lo crea con las hojas Reseñas, Resumen y Cómo usar.
- No duplica: usa el ID de la reseña. Si una reseña cambió (estrellas, texto o respuesta), la actualiza.
- No toca las columnas amarillas (las completa la familia).
- Imprime un resumen en JSON con las reseñas nuevas y las negativas nuevas (3 estrellas o menos) para avisar.
"""
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

xlsx, origen, dia = sys.argv[1], sys.argv[2], sys.argv[3]
completa = len(sys.argv) > 4 and sys.argv[4] == "completa"
momento = datetime.fromisoformat(dia if "T" in dia else dia + "T12:00")
hoy = momento.date()
resenas = json.load(open(origen, encoding="utf-8"))

AZUL = "1C2754"
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
CABEZA = PatternFill("solid", fgColor=AZUL)
LINEA = Border(bottom=Side(style="thin", color="E7CFD8"))
F = lambda **k: Font(name="Arial", **k)
COLUMNAS = [
    ("N°", 6), ("Vista por primera vez", 13), ("Fecha aprox.", 12), ("Google dice", 14), ("Nombre en Google", 26),
    ("Estrellas", 9), ("Comentario", 60), ("Temas", 26), ("¿Respondida?", 12), ("Respuesta de la tienda", 40),
    ("Perfil en Google", 16), ("¿Familiar o trabaja en la tienda?", 16), ("Notas", 30), ("Estado", 14), ("ID de la reseña", 18),
]
TEMAS = [
    ("atención", r"atenci|atend|amabl|cari[ñn]|calidez|c[aá]lid|humana|personaliz|vecina|due[ñn]a|profesional|dispuest|preocupaci"),
    ("variedad", r"variedad|todo lo que|de todo|todo en un|encuentr|surtid|alternativ"),
    ("calidad", r"calidad|buen(os|as)? producto|excelente(s)? producto|prendas"),
    ("precio", r"precio|barat|econ[oó]mic|accesible|bolsillo|alcance"),
    ("recomienda", r"recomiend|recomendad|10 de 10|100|1000"),
    ("encargos", r"encarg|avisan|pedido"),
]


def fecha_aprox(texto):
    t = (texto or "").lower().replace("hace ", "").strip()
    if not t:
        return None
    n = 1 if re.match(r"^(un|una)\b", t) else int(re.match(r"^(\d+)", t).group(1)) if re.match(r"^\d+", t) else 1
    if "minuto" in t or "segundo" in t:
        return hoy
    if "hora" in t:
        return (momento - timedelta(hours=n)).date()
    if "día" in t or "dia" in t:
        return hoy - timedelta(days=n)
    if "semana" in t:
        return hoy - timedelta(weeks=n)
    if "mes" in t:
        return hoy - timedelta(days=30 * n)
    if "año" in t or "ano" in t:
        return hoy - timedelta(days=365 * n)
    return None


def temas_de(texto):
    t = (texto or "").lower()
    return ", ".join(nombre for nombre, patron in TEMAS if re.search(patron, t))


def crear_libro(ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reseñas"
    for i, (titulo, ancho) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = CABEZA
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:O1"
    dv = DataValidation(type="list", formula1='"Sí,No,No sé"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("L2:L2000")

    rs = wb.create_sheet("Resumen")
    rs.column_dimensions["A"].width = 46
    rs.column_dimensions["B"].width = 16
    filas = [
        ("Reseñas de Google de Angie's Secret", None),
        (None, None),
        ("Reseñas registradas", "=COUNTA(Reseñas!E2:E2000)"),
        ("Visibles hoy en Google", '=COUNTIF(Reseñas!N2:N2000,"Visible")'),
        ("Promedio de estrellas", '=IFERROR(AVERAGE(Reseñas!F2:F2000),0)'),
        ("5 estrellas", '=COUNTIF(Reseñas!F2:F2000,5)'),
        ("4 estrellas", '=COUNTIF(Reseñas!F2:F2000,4)'),
        ("3 estrellas", '=COUNTIF(Reseñas!F2:F2000,3)'),
        ("2 estrellas", '=COUNTIF(Reseñas!F2:F2000,2)'),
        ("1 estrella", '=COUNTIF(Reseñas!F2:F2000,1)'),
        ("Negativas o neutras (3 estrellas o menos)", '=COUNTIF(Reseñas!F2:F2000,"<=3")'),
        ("Con comentario escrito", '=COUNTIF(Reseñas!G2:G2000,"?*")'),
        ("Respondidas por la tienda", '=COUNTIF(Reseñas!I2:I2000,"Sí")'),
        ("% respondidas", '=IFERROR(B13/B3,0)'),
        ("Nuevas en los últimos 7 días", '=COUNTIF(Reseñas!B2:B2000,">="&(TODAY()-7))'),
        ("Marcadas como familiar o de la tienda", '=COUNTIF(Reseñas!L2:L2000,"Sí")'),
        (None, None),
        ("Lo que más destacan (reseñas que lo mencionan)", None),
    ] + [(f"  {n}", f'=COUNTIF(Reseñas!H2:H2000,"*{n}*")') for n, _ in TEMAS] + [
        (None, None),
        ("Última actualización", None),
    ]
    for i, (a, b) in enumerate(filas, start=1):
        if a is not None:
            rs.cell(row=i, column=1, value=a).font = F(bold=(i in (1, 18)), size=(14 if i == 1 else 11), color=AZUL if i == 1 else None)
        if b is not None:
            c = rs.cell(row=i, column=2, value=b)
            c.font = F()
            if a == "Promedio de estrellas":
                c.number_format = "0.0"
            if a == "% respondidas":
                c.number_format = "0%"
    rs["A{}".format(len(filas))].font = F(bold=True)

    cu = wb.create_sheet("Cómo usar")
    cu.column_dimensions["A"].width = 110
    texto = [
        "Cómo usar este Excel",
        "",
        "Qué es: el registro de las reseñas que la gente deja en el perfil de Angie's Secret en Google, con su nombre tal como aparece en Google.",
        "Se actualiza solo: una tarea programada lee las reseñas en Google Maps y agrega aquí las nuevas, sin duplicar ni borrar nada.",
        "Para que la tarea funcione, este computador debe estar encendido con la app de Claude abierta a la hora programada.",
        "",
        "Columnas blancas: las llena la actualización automática. No hace falta tocarlas.",
        "Columnas amarillas: las llena la familia.",
        "  ¿Familiar o trabaja en la tienda?: Google no permite reseñas de familiares ni de quienes trabajan en el negocio y puede borrarlas.",
        "     Marcar 'Sí' ayuda a saber cuántas podrían desaparecer y a no pedir más reseñas así.",
        "  Notas: lo que la familia quiera anotar (por ejemplo, si ya se le agradeció a la persona).",
        "Estado: 'Ya no aparece' significa que la reseña estaba antes y ahora Google no la muestra (la borró la persona o la filtró Google).",
        "Fecha aprox.: Google solo dice 'hace 3 días' o 'hace un mes'; la fecha se calcula desde el día en que se leyó.",
        "",
        "Privacidad: los nombres son públicos en Google, pero este Excel es interno. No lo compartas ni publiques su contenido.",
        "Para responder reseñas: Perfil de Empresa en Google > Leer opiniones > Responder. Responder a todas mejora el perfil.",
    ]
    for i, linea in enumerate(texto, start=1):
        c = cu.cell(row=i, column=1, value=linea)
        c.font = F(bold=(i == 1), size=(14 if i == 1 else 11), color=AZUL if i == 1 else None)
        c.alignment = Alignment(wrap_text=True)
    wb.save(ruta)


if not Path(xlsx).exists():
    crear_libro(xlsx)

wb = load_workbook(xlsx)
ws = wb["Reseñas"]
filas_por_id, ultima = {}, 1
for r in range(2, ws.max_row + 1):
    if ws.cell(row=r, column=15).value:
        filas_por_id[str(ws.cell(row=r, column=15).value)] = r
        ultima = r

nuevas, negativas, actualizadas = [], [], 0
vistos = set()
# Google las entrega de la más nueva a la más antigua: se agregan de la más antigua a la más nueva.
for res in reversed(resenas):
    rid = str(res.get("id") or "").strip()
    if not rid:
        continue
    vistos.add(rid)
    estrellas = res.get("estrellas")
    texto = (res.get("texto") or "").strip()
    respuesta = (res.get("respuesta") or "").strip()
    if rid in filas_por_id:
        r = filas_por_id[rid]
        cambio = False
        for col, valor in ((6, estrellas), (7, texto or None), (10, respuesta or None)):
            if valor is not None and ws.cell(row=r, column=col).value != valor:
                ws.cell(row=r, column=col, value=valor)
                cambio = True
        ws.cell(row=r, column=8, value=temas_de(texto) or None)
        ws.cell(row=r, column=9, value="Sí" if respuesta or ws.cell(row=r, column=10).value else "No")
        ws.cell(row=r, column=14, value="Visible")
        actualizadas += int(cambio)
        continue
    ultima += 1
    r = ultima
    valores = {
        1: f"=ROW()-1", 2: hoy, 3: fecha_aprox(res.get("fecha")), 4: res.get("fecha") or None, 5: (res.get("nombre") or "").strip(),
        6: estrellas, 7: texto or None, 8: temas_de(texto) or None, 9: "Sí" if respuesta else "No", 10: respuesta or None,
        11: res.get("perfil") or None, 14: "Visible", 15: rid,
    }
    for col in range(1, 16):
        c = ws.cell(row=r, column=col, value=valores.get(col))
        c.font = F(color="808080", size=8) if col == 15 else F()
        c.alignment = Alignment(wrap_text=col in (7, 10), vertical="top")
        c.border = LINEA
        if col in (12, 13):
            c.fill = AMARILLO
    for col in (2, 3):
        ws.cell(row=r, column=col).number_format = "DD-MM-YYYY"
    filas_por_id[rid] = r
    item = {"nombre": valores[5], "estrellas": estrellas, "texto": texto[:300], "fecha": res.get("fecha")}
    nuevas.append(item)
    if isinstance(estrellas, int) and estrellas <= 3:
        negativas.append(item)

ya_no = 0
if completa:
    for rid, r in filas_por_id.items():
        if rid not in vistos and ws.cell(row=r, column=14).value != "Ya no aparece":
            ws.cell(row=r, column=14, value="Ya no aparece").font = F(color="C00000", bold=True)
            ya_no += 1

rs = wb["Resumen"]
for i in range(1, rs.max_row + 1):
    if rs.cell(row=i, column=1).value == "Última actualización":
        c = rs.cell(row=i, column=2, value=hoy)
        c.number_format = "DD-MM-YYYY"
        c.font = F()
wb.calculation.fullCalcOnLoad = True
wb.save(xlsx)
print(json.dumps({"nuevas": len(nuevas), "negativas_nuevas": negativas, "actualizadas": actualizadas,
                  "ya_no_aparecen": ya_no, "total": len(filas_por_id), "detalle_nuevas": nuevas[:20]}, ensure_ascii=False))
